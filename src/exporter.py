"""In-memory CSV, JSON and styled Excel exports without sensitive credentials."""

from __future__ import annotations

import json
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


EXPORT_COLUMNS = [
    "time",
    "pool_type",
    "name",
    "item_type",
    "rank_type",
    "pity_count",
    "is_off_banner",
    "is_guaranteed",
    "five_star_result_type",
    "id",
    "uid",
    "gacha_id",
    "gacha_type",
    "item_id",
]


def export_frame(frame: pd.DataFrame) -> pd.DataFrame:
    output = frame.copy()
    for column in EXPORT_COLUMNS:
        if column not in output.columns:
            output[column] = None
    output = output[EXPORT_COLUMNS]
    output["time"] = pd.to_datetime(output["time"], errors="coerce")
    return output


def to_csv_bytes(frame: pd.DataFrame) -> bytes:
    output = export_frame(frame).copy()
    output["time"] = output["time"].dt.strftime("%Y-%m-%d %H:%M:%S")
    return output.to_csv(index=False).encode("utf-8-sig")


def to_json_bytes(frame: pd.DataFrame) -> bytes:
    output = export_frame(frame).copy()
    output["time"] = output["time"].dt.strftime("%Y-%m-%d %H:%M:%S")
    records = json.loads(output.to_json(orient="records", force_ascii=False))
    payload = {"schema_version": 1, "records": records}
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def _summary_rows(summary: dict[str, Any]) -> list[tuple[str, Any]]:
    character = summary["pools"]["character_event"]
    lightcone = summary["pools"]["lightcone_event"]
    return [
        ("Total Pulls", summary["total_pulls"]),
        ("Five Stars", summary["five_stars"]),
        ("Average Pity", summary["average_pity"]),
        ("Character 50/50 Win Rate", character["small_pity_win_rate"]),
        ("Character Off-banner Count", character["off_banner_count"]),
        ("Light Cone Featured Win Rate", lightcone["small_pity_win_rate"]),
        ("Current Character Pity", character["current_pity"]),
        ("Character Guaranteed Next", character["guaranteed_next"]),
        ("Current Light Cone Pity", lightcone["current_pity"]),
        ("Light Cone Guaranteed Next", lightcone["guaranteed_next"]),
    ]


def to_excel_bytes(frame: pd.DataFrame, summary: dict[str, Any]) -> bytes:
    output = export_frame(frame)
    five_stars = output[output["rank_type"].eq(5)].copy()
    summary_frame = pd.DataFrame(_summary_rows(summary), columns=["Metric", "Value"])
    buffer = BytesIO()

    with pd.ExcelWriter(buffer, engine="openpyxl", datetime_format="yyyy-mm-dd hh:mm:ss") as writer:
        summary_frame.to_excel(writer, sheet_name="Summary", index=False)
        output.to_excel(writer, sheet_name="Pull History", index=False)
        five_stars.to_excel(writer, sheet_name="Five Star History", index=False)

        workbook = writer.book
        header_fill = PatternFill("solid", fgColor="182A4A")
        header_font = Font(color="FFFFFF", bold=True)
        accent_fill = PatternFill("solid", fgColor="263F70")

        for worksheet in workbook.worksheets:
            worksheet.sheet_view.showGridLines = False
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for column_cells in worksheet.columns:
                values = [str(cell.value) if cell.value is not None else "" for cell in column_cells]
                width = min(max(max(map(len, values), default=0) + 2, 12), 34)
                worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width

        summary_sheet = workbook["Summary"]
        summary_sheet.column_dimensions["A"].width = 34
        summary_sheet.column_dimensions["B"].width = 20
        for row in range(2, summary_sheet.max_row + 1):
            summary_sheet.cell(row, 1).fill = accent_fill
            summary_sheet.cell(row, 1).font = Font(color="FFFFFF", bold=True)
        summary_sheet["B5"].number_format = "0.0%"
        summary_sheet["B7"].number_format = "0.0%"

    buffer.seek(0)
    return buffer.getvalue()

