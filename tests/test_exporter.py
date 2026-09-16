from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from src.analyzer import analyze_account
from src.exporter import to_csv_bytes, to_excel_bytes, to_json_bytes
from src.parser import records_to_dataframe
from src.pity import apply_pity_and_guarantee


def test_exports_have_expected_sheets_and_no_authkey() -> None:
    record = {
        "id": "1",
        "uid": "800000001",
        "gacha_id": "banner",
        "gacha_type": "11",
        "item_id": "item",
        "name": "阮·梅",
        "item_type": "角色",
        "rank_type": "5",
        "time": "2025-01-01 12:00:00",
        "authkey": "must-not-export",
    }
    frame = apply_pity_and_guarantee(records_to_dataframe([record]))
    summary = analyze_account(frame)

    csv_bytes = to_csv_bytes(frame)
    json_bytes = to_json_bytes(frame)
    excel_bytes = to_excel_bytes(frame, summary)
    assert b"must-not-export" not in csv_bytes
    assert b"must-not-export" not in json_bytes

    workbook = load_workbook(BytesIO(excel_bytes), read_only=True)
    assert workbook.sheetnames == ["Summary", "Pull History", "Five Star History"]

